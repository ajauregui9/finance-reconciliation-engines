"""Pruebas de la salida del cierre: JSON determinista, xlsx numérico verificado y CLI."""
from __future__ import annotations

import dataclasses
import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

import test_amazon_soporte as soporte
from cierre_amazon.cierre import ResultadoCorrida, cerrar
from cierre_amazon.cli import main
from cierre_amazon.config import cargar_config
from cierre_amazon.insumos import cargar_insumos
from cierre_amazon.salida_json import a_dict, a_json
from cierre_amazon.salida_xlsx import escribir_xlsx, verificar_xlsx

HUSO = "UTC-07:00"
ABRIL = soporte.csv_liquidacion("2026-04", HUSO, [
    soporte.venta("2026-04-10 10:00:00", "P-1", "1000.00", fecha_pedido="2026-04-09 10:00:00",
                  comision="-150.00", logistica="-100.00", retencion="-80.00"),
    soporte.cargo("2026-04-09 11:00:00", "Publicidad", "-116.00"),
    soporte.cargo("2026-04-05 10:00:00", "Almacenaje", "-300.00"),
])
MAYO = soporte.csv_liquidacion("2026-05", HUSO, [
    soporte.venta("2026-05-03 10:00:00", "P-2", "500.00", fecha_pedido="2026-04-29 10:00:00",
                  comision="-80.00", logistica="-45.00", retencion="-60.00"),
    soporte.venta("2026-05-12 10:00:00", "P-3", "2000.00", fecha_pedido="2026-05-11 10:00:00",
                  comision="-300.00", logistica="-160.00", retencion="-300.00"),
])
CONSOLA = "Fecha,Campaña,Gasto\n2026-04-02,NOR | Busqueda,100.00\n"
FACTURAS = [{"mes": "2026-04", "folio": "PUB-1", "subtotal": "100.00", "iva": "16.00", "total": "116.00"},
            {"mes": "2026-05", "folio": "PUB-2", "subtotal": "0.00", "iva": "0.00", "total": "0.00"}]


@pytest.fixture()
def datos(tmp_path: Path) -> Path:
    return soporte.escribir_insumos(tmp_path / "datos", {"2026-04": ABRIL, "2026-05": MAYO}, consola=CONSOLA,
                                    facturas=FACTURAS, venta_erp={"2026-04": "1500.00", "2026-05": "2000.00"})


def _resultado(directorio: Path) -> ResultadoCorrida:
    config = cargar_config(soporte.CONFIG_EJEMPLO)
    return cerrar(cargar_insumos(directorio, config), config)


def _mes(documento: Dict[str, Any], mes: str) -> Dict[str, Any]:
    return next(m for m in documento["meses"] if m["mes"] == mes)


class TestJson:
    def test_montos_como_texto_y_nulos_explicitos(self, datos: Path) -> None:
        documento = a_dict(_resultado(datos))
        assert documento["esquema"] == "cierre_amazon/v1"
        abril = _mes(documento, "2026-04")
        assert abril["gasto_devengado"]["total"] == "775.00"
        assert abril["gasto_devengado"]["publicidad"] == "100.00"
        assert abril["gasto_devengado"]["tarifas"]["comision"] == {
            "cohorte_mes": "150.00", "provision": "75.00", "cola_real": "0.00", "reverso": "0.00", "total": "225.00"}
        assert abril["devengo"]["tasas_efectivas"]["comision"] == "0.150000"
        assert abril["devengo"]["provision_anterior"] is None
        assert abril["retencion"]["tasa_observada"] == "0.080000"
        assert abril["publicidad"]["consola"] == {"filas": 1, "sin_iva": "100.00"}
        assert documento["archivos"][0] == {"archivo": "liquidacion_2026-04.csv", "filas": 3, "huso": HUSO,
                                            "periodo": "2026-04", "total": "254.00"}

    def test_alerta_de_retencion_aparece_en_el_mes(self, datos: Path) -> None:
        mayo = _mes(a_dict(_resultado(datos)), "2026-05")
        alertas = [a for a in mayo["alertas"] if a["codigo"] == "retencion_fuera_de_banda"]
        assert len(alertas) == 1
        assert alertas[0]["datos"]["tasa_observada"] == "0.144000"
        assert mayo["retencion"]["dentro_de_banda"] is False
        assert all(c["ok"] for c in mayo["candados"])

    def test_dos_corridas_dan_los_mismos_bytes(self, datos: Path) -> None:
        primero = a_json(_resultado(datos)).encode("utf-8")
        segundo = a_json(_resultado(datos)).encode("utf-8")
        assert hashlib.sha256(primero).hexdigest() == hashlib.sha256(segundo).hexdigest()
        assert primero.endswith(b"\n")

    def test_sin_rutas_absolutas_ni_floats(self, datos: Path) -> None:
        texto = a_json(_resultado(datos))
        assert str(datos) not in texto

        def recorrer(valor: Any) -> None:
            assert not isinstance(valor, float)
            if isinstance(valor, dict):
                for v in valor.values():
                    recorrer(v)
            elif isinstance(valor, list):
                for v in valor:
                    recorrer(v)

        recorrer(json.loads(texto))

    def test_filtrar_meses(self, datos: Path) -> None:
        documento = a_dict(_resultado(datos), meses=["2026-05"])
        assert [m["mes"] for m in documento["meses"]] == ["2026-05"]
        with pytest.raises(ValueError, match="2026-09"):
            a_dict(_resultado(datos), meses=["2026-09"])


class TestXlsx:
    def test_montos_numericos_y_verificacion_limpia(self, datos: Path, tmp_path: Path) -> None:
        resultado = _resultado(datos)
        ruta = tmp_path / "cierre.xlsx"
        escribir_xlsx(resultado, ruta)
        libro = load_workbook(ruta, data_only=True)
        assert libro.sheetnames == ["Resumen", "Gasto devengado", "Alertas", "Candados"]
        hoja = libro["Resumen"]
        encabezado = [c.value for c in hoja[1]]
        columna = encabezado.index("Gasto devengado") + 1
        celda = hoja.cell(row=2, column=columna)
        assert isinstance(celda.value, (int, float))
        assert celda.number_format == "#,##0.00"
        assert verificar_xlsx(ruta, resultado) == ()

    def test_verificacion_detecta_un_monto_alterado(self, datos: Path, tmp_path: Path) -> None:
        resultado = _resultado(datos)
        ruta = tmp_path / "cierre.xlsx"
        escribir_xlsx(resultado, ruta)
        libro = load_workbook(ruta)
        hoja = libro["Gasto devengado"]
        hoja.cell(row=2, column=4).value = 1.23
        libro["Resumen"].cell(row=3, column=4).value = "2000.00"
        libro.save(ruta)
        diferencias = verificar_xlsx(ruta, resultado)
        assert len(diferencias) == 2
        assert any("texto" in d for d in diferencias)

    def test_verificacion_detecta_hoja_faltante(self, datos: Path, tmp_path: Path) -> None:
        resultado = _resultado(datos)
        ruta = tmp_path / "cierre.xlsx"
        escribir_xlsx(resultado, ruta)
        libro = load_workbook(ruta)
        del libro["Resumen"]
        libro.save(ruta)
        assert any("Resumen" in d for d in verificar_xlsx(ruta, resultado))


class TestXlsxReproducible:
    """El libro no depende del reloj: ni las entradas del zip ni las propiedades del documento."""

    def test_dos_escrituras_dan_el_mismo_sha256(self, datos: Path, tmp_path: Path) -> None:
        resultado = _resultado(datos)
        primero, segundo = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
        escribir_xlsx(resultado, primero)
        escribir_xlsx(resultado, segundo)
        assert hashlib.sha256(primero.read_bytes()).hexdigest() == hashlib.sha256(segundo.read_bytes()).hexdigest()
        assert verificar_xlsx(primero, resultado) == ()

    def test_entradas_del_zip_con_fecha_y_permisos_fijos(self, datos: Path, tmp_path: Path) -> None:
        ruta = tmp_path / "cierre.xlsx"
        escribir_xlsx(_resultado(datos), ruta)
        with ZipFile(ruta) as archivo:
            entradas = archivo.infolist()
        assert entradas and {e.date_time for e in entradas} == {(1980, 1, 1, 0, 0, 0)}
        assert {(e.create_system, e.external_attr >> 16) for e in entradas} == {(3, 0o100644)}

    def test_propiedades_fijas_al_ultimo_dia_del_ultimo_mes_del_libro(self, datos: Path, tmp_path: Path) -> None:
        resultado = _resultado(datos)
        ruta, solo_abril = tmp_path / "cierre.xlsx", tmp_path / "abril.xlsx"
        escribir_xlsx(resultado, ruta)
        escribir_xlsx(resultado, solo_abril, ["2026-04"])
        propiedades = load_workbook(ruta).properties
        assert (propiedades.creator, propiedades.lastModifiedBy) == ("cierre_amazon", "cierre_amazon")
        assert propiedades.created == propiedades.modified == datetime(2026, 5, 31)
        assert load_workbook(solo_abril).properties.modified == datetime(2026, 4, 30)

    def test_libro_sin_meses_falla_sin_escribir(self, datos: Path, tmp_path: Path) -> None:
        vacio = dataclasses.replace(_resultado(datos), meses=())
        with pytest.raises(ValueError, match="sin meses"):
            escribir_xlsx(vacio, tmp_path / "vacio.xlsx")
        assert not (tmp_path / "vacio.xlsx").exists()


class TestCli:
    def _correr(self, *argumentos: str) -> int:
        return main(list(argumentos))

    def test_escribe_json_y_xlsx(self, datos: Path, tmp_path: Path) -> None:
        salida, libro = tmp_path / "cierre.json", tmp_path / "cierre.xlsx"
        codigo = self._correr("--datos", str(datos), "--salida", str(salida), "--xlsx", str(libro))
        assert codigo == 0
        assert json.loads(salida.read_text(encoding="utf-8"))["ok"] is True
        assert libro.is_file()

    def test_json_a_stdout_y_resumen_a_stderr(self, datos: Path, capsys: pytest.CaptureFixture[str]) -> None:
        assert self._correr("--datos", str(datos), "--mes", "2026-04") == 0
        capturado = capsys.readouterr()
        assert [m["mes"] for m in json.loads(capturado.out)["meses"]] == ["2026-04"]
        assert "2026-04" in capturado.err

    def test_fuente_de_publicidad_desde_la_linea_de_comandos(self, datos: Path,
                                                            capsys: pytest.CaptureFixture[str]) -> None:
        assert self._correr("--datos", str(datos), "--fuente-publicidad", "factura") == 0
        documento = json.loads(capsys.readouterr().out)
        assert documento["parametros"]["fuente_publicidad"] == "factura"

    @pytest.mark.parametrize("extra", [["--mes", "2026-09"], ["--mes", "abril"], ["--fuente-publicidad", "radio"]])
    def test_argumentos_invalidos_salen_con_2(self, datos: Path, extra: List[str],
                                              capsys: pytest.CaptureFixture[str]) -> None:
        assert self._correr("--datos", str(datos), *extra) == 2
        assert capsys.readouterr().out == ""

    def test_insumo_o_config_invalidos_salen_con_2(self, tmp_path: Path,
                                                    capsys: pytest.CaptureFixture[str]) -> None:
        assert self._correr("--datos", str(tmp_path / "nada")) == 2
        assert "no existe" in capsys.readouterr().err
        datos = soporte.escribir_insumos(tmp_path / "d", {"2026-04": ABRIL})
        assert self._correr("--datos", str(datos), "--config", str(tmp_path / "falta.json")) == 2

    def test_candado_en_falla_sale_con_1(self, tmp_path: Path) -> None:
        texto = soporte.csv_liquidacion("2026-04", HUSO, [soporte.cargo("2026-04-12 10:00:00", "Cargo nuevo", "-1.00")])
        datos = soporte.escribir_insumos(tmp_path / "d", {"2026-04": texto}, consola="Fecha,Campaña,Gasto\n")
        salida = tmp_path / "c.json"
        assert self._correr("--datos", str(datos), "--salida", str(salida)) == 1
        assert json.loads(salida.read_text(encoding="utf-8"))["ok"] is False
