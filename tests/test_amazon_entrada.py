"""Contrato de la frontera de entrada del cierre de Amazon, en cada insumo.

Montos: texto con comas de miles opcionales y bien colocadas, o número JSON; hasta dos
decimales, dígitos ASCII. En la liquidación una celda vacía es cero (spec 2.1); en las
demás fuentes es error. Fechas y husos con dígitos ASCII y forma fija. Un insumo fuera de
contrato termina la CLI con código 2.
"""
from __future__ import annotations

import json
from datetime import timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict

import pytest
from openpyxl import load_workbook

import test_amazon_soporte as soporte
from cierre_amazon import dinero, huso
from cierre_amazon.cierre import cerrar
from cierre_amazon.cli import main
from cierre_amazon.config import cargar_config
from cierre_amazon.errores import ErrorInsumo
from cierre_amazon.insumos import cargar_insumos
from cierre_amazon.lector_fuentes import facturas_desde_dict, leer_consola_texto, venta_erp_desde_dict
from cierre_amazon.lector_liquidacion import leer_liquidacion_texto
from cierre_amazon.salida_xlsx import escribir_xlsx, verificar_xlsx

CONSOLA = "Fecha,Campaña,Gasto\n2026-04-10,X,1.00\n"
FUERA_DE_CONTRATO = ["١٠٠.٠٠", "100.005", "1 000.00", "1,00.00", "1,2,3", "1e2", ".5", 0.1 + 0.2, 100.005]


def _leer(texto: str) -> Any:
    return leer_liquidacion_texto(texto, "liquidacion_2026-04.csv", "MXN")


def _venta(**cambios: str) -> str:
    fila = soporte.venta("2026-04-10 10:00:00", "P-1", "100.00")
    fila.update(cambios)
    return soporte.csv_liquidacion("2026-04", "UTC-07:00", [fila])


# ------------------------------------------------------------------ a_decimal
@pytest.mark.parametrize("valor, esperado", [("1,234.56", "1234.56"), (" 7 ", "7.00"), (7, "7.00"),
                                             (100.5, "100.50"), (Decimal("2.5"), "2.50")])
def test_monto_valido(valor: object, esperado: str) -> None:
    assert str(dinero.a_decimal(valor)) == esperado


@pytest.mark.parametrize("valor", FUERA_DE_CONTRATO + ["", "   "])
def test_monto_fuera_de_contrato(valor: object) -> None:
    with pytest.raises(ValueError, match="monto"):
        dinero.a_decimal(valor)


def test_vacio_es_cero_solo_cuando_se_pide() -> None:
    assert dinero.a_decimal("  ", vacio_es_cero=True) == Decimal("0.00")
    with pytest.raises(ValueError, match="monto"):
        dinero.a_decimal(None, vacio_es_cero=True)


@pytest.mark.parametrize("valor", ["٠.١٦", "0,16", "0.16 0", 0.16, 1])
def test_tasa_solo_texto_ascii(valor: object) -> None:
    with pytest.raises(ValueError):
        dinero.a_tasa(valor)


def test_tasa_releida_del_xlsx() -> None:
    assert dinero.tasa_releida(0.0812345678912345) == Decimal("0.0812345678912345")
    assert dinero.tasa_releida(1) == Decimal(1)
    for valor in (float("nan"), True, "0.08"):
        with pytest.raises(ValueError):
            dinero.tasa_releida(valor)


# ----------------------------------------------------------------- liquidación
@pytest.mark.parametrize("monto", ["١٠٠.٠٠", "100.005", "1 00.00", "1,00.00"])
def test_liquidacion_con_monto_fuera_de_contrato(monto: str) -> None:
    with pytest.raises(ErrorInsumo, match="Venta de productos"):
        _leer(_venta(venta_productos=monto, total=monto))


def test_liquidacion_celda_vacia_sigue_siendo_cero() -> None:
    (fila,) = _leer(_venta(envio_cobrado="")).filas
    assert fila.montos.envio_cobrado == Decimal("0.00")


@pytest.mark.parametrize("cantidad", ["١", "1.5", "1_0", "+1"])
def test_liquidacion_cantidad_entero_ascii(cantidad: str) -> None:
    with pytest.raises(ErrorInsumo, match="Cantidad"):
        _leer(_venta(cantidad=cantidad))


def test_liquidacion_cantidad_negativa_y_vacia() -> None:
    assert _leer(_venta(cantidad="-2")).filas[0].cantidad == -2
    assert _leer(_venta(cantidad=" ")).filas[0].cantidad == 0


@pytest.mark.parametrize("columna, valor", [("fecha_hora", "٢٠٢٦-٠٤-١٠ ١٠:٠٠:٠٠"), ("fecha_hora", "2026-4-10 10:0:0"),
                                            ("fecha_pedido", "2026-04-10T10:00:00")])
def test_liquidacion_fechas_forma_fija(columna: str, valor: str) -> None:
    with pytest.raises(ErrorInsumo, match="AAAA-MM-DD HH:MM:SS"):
        _leer(_venta(**{columna: valor}))


@pytest.mark.parametrize("periodo, huso_texto", [("٢٠٢٦-٠٤", "UTC-07:00"), ("2026-04", "UTC-٠٧:٠٠")])
def test_liquidacion_preambulo_ascii(periodo: str, huso_texto: str) -> None:
    with pytest.raises(ErrorInsumo):
        _leer(soporte.csv_liquidacion(periodo, huso_texto, []))


def test_huso_y_fecha_hora_delegan_en_la_frontera() -> None:
    assert huso.leer_fecha_hora(" 2026-04-10 10:00:00 ", timezone.utc).hour == 10
    for malo in ("UTC-٠٧:٠٠", "UTC-7:00"):
        with pytest.raises(ValueError, match="UTC±HH:MM"):
            huso.interpretar_huso(malo)
    with pytest.raises(ValueError, match="AAAA-MM"):
        huso.validar_mes("٢٠٢٦-٠٤")


# --------------------------------------------------------------------- fuentes
@pytest.mark.parametrize("gasto", ["١٠.٠٠", "10.005", "1,0.00", ""])
def test_consola_gasto_fuera_de_contrato(gasto: str) -> None:
    with pytest.raises(ErrorInsumo, match="Gasto"):
        leer_consola_texto(f"Fecha,Campaña,Gasto\n2026-04-01,X,\"{gasto}\"\n", "publicidad_consola.csv")


@pytest.mark.parametrize("fecha", ["2026-4-1", "٢٠٢٦-٠٤-٠١"])
def test_consola_fecha_forma_fija(fecha: str) -> None:
    with pytest.raises(ErrorInsumo, match="Fecha"):
        leer_consola_texto(f"Fecha,Campaña,Gasto\n{fecha},X,1.00\n", "publicidad_consola.csv")


@pytest.mark.parametrize("subtotal", ["100.005", "١٠٠", 0.1 + 0.2, ""])
def test_facturas_monto_fuera_de_contrato(subtotal: object) -> None:
    documento: Dict[str, Any] = {"esquema": "facturas_publicidad/v1", "facturas": [
        {"mes": "2026-04", "folio": "F-1", "subtotal": subtotal, "iva": "16.00", "total": "116.00"}]}
    with pytest.raises(ErrorInsumo, match="subtotal"):
        facturas_desde_dict(documento, "facturas_publicidad.json")


def test_facturas_aceptan_numero_json_exacto() -> None:
    documento: Dict[str, Any] = {"esquema": "facturas_publicidad/v1", "facturas": [
        {"mes": "2026-04", "folio": "F-1", "subtotal": 100.5, "iva": 16, "total": "116.50"}]}
    assert facturas_desde_dict(documento, "facturas_publicidad.json")[0].subtotal == Decimal("100.50")


@pytest.mark.parametrize("valor", ["1,000.005", "١٠٠٠", 0.1 + 0.2])
def test_venta_erp_fuera_de_contrato(valor: object) -> None:
    with pytest.raises(ErrorInsumo, match="monto"):
        venta_erp_desde_dict({"esquema": "venta_erp/v1", "venta_neta_por_mes_de_pedido": {"2026-04": valor}},
                             "venta_erp.json")


# ------------------------------------------------------------------- xlsx y CLI
def test_xlsx_con_float_de_tres_decimales_es_diferencia_no_excepcion(tmp_path: Path) -> None:
    config = cargar_config(soporte.CONFIG_EJEMPLO)
    resultado = cerrar(cargar_insumos(soporte.DIR_SINTETICO, config), config)
    ruta = tmp_path / "cierre.xlsx"
    escribir_xlsx(resultado, ruta)
    libro = load_workbook(ruta)
    libro["Gasto devengado"].cell(row=2, column=4).value = 1.005
    libro.save(ruta)
    (diferencia,) = verificar_xlsx(ruta, resultado)
    assert "Gasto devengado!D2" in diferencia and "decimales" in diferencia


def test_cli_insumo_fuera_de_contrato_sale_con_2(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    datos = soporte.escribir_insumos(tmp_path / "datos", {"2026-04": _venta(venta_productos="100.005",
                                                                            total="100.005")}, consola=CONSOLA)
    salida = tmp_path / "cierre.json"
    assert main(["--datos", str(datos), "--salida", str(salida)]) == 2
    assert not salida.exists()
    assert "a lo más 2 decimales" in capsys.readouterr().err


def test_cli_venta_erp_con_float_ambiguo_sale_con_2(tmp_path: Path) -> None:
    datos = soporte.escribir_insumos(tmp_path / "datos", {"2026-04": _venta()}, consola=CONSOLA)
    (datos / "venta_erp.json").write_text(json.dumps(
        {"esquema": "venta_erp/v1", "venta_neta_por_mes_de_pedido": {"2026-04": 0.1 + 0.2}}), encoding="utf-8")
    assert main(["--datos", str(datos), "--salida", str(tmp_path / "cierre.json")]) == 2
