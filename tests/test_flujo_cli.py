"""Pruebas de integración y punta a punta del CLI `python3 -m flujo`.

Corren sobre los insumos versionados de examples/synthetic_flujo/ y la config de ejemplo.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import zipfile
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Dict, List

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
MODULO = RAIZ / "skills" / "cashflow-forecasting"
if str(MODULO) not in sys.path:
    sys.path.insert(0, str(MODULO))

from flujo.cli import main  # noqa: E402

CONFIG_EJEMPLO = MODULO / "config" / "flujo_ejemplo.json"
DATOS = RAIZ / "examples" / "synthetic_flujo"
CORTE = "2026-11-02"
SALIDAS = ("filas_contrato.jsonl", "flujo_13_semanas.json", "flujo_13_semanas.xlsx")
MONTO_TXT = re.compile(r"^-?\d+\.\d{2}$")
D = Decimal
SEGUNDOS_POR_DIA = 86_400
FECHA_ZIP = (1980, 1, 1, 0, 0, 0)


def _correr(salida: Path, *extra: str, config: Path = CONFIG_EJEMPLO, corte: str = CORTE) -> int:
    return main(["--config", str(config), "--corte", corte, "--salida", str(salida), *extra])


def _sha(ruta: Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def _json(salida: Path) -> Dict[str, Any]:
    datos: Dict[str, Any] = json.loads((salida / "flujo_13_semanas.json").read_text(encoding="utf-8"))
    return datos


def _copiar_config(tmp_path: Path, **cambios: Any) -> Path:
    """Copia la config de ejemplo con rutas absolutas y aplica cambios de primer nivel."""
    datos = json.loads(CONFIG_EJEMPLO.read_text(encoding="utf-8"))
    datos["insumos"] = {k: str((CONFIG_EJEMPLO.parent / v).resolve()) for k, v in datos["insumos"].items()}
    datos.update(cambios)
    ruta = tmp_path / "config.json"
    ruta.write_text(json.dumps(datos), encoding="utf-8")
    return ruta


# --- corrida exitosa --------------------------------------------------------------------


def test_corrida_escribe_las_tres_salidas(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert _correr(tmp_path / "out") == 0
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == list(SALIDAS)
    resumen = json.loads(capsys.readouterr().out)
    assert resumen["semanas"] == 13
    assert resumen["archivos"] == list(SALIDAS)


def test_json_tiene_13_semanas_encadenadas_y_montos_como_texto(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    datos = _json(tmp_path)
    assert datos["esquema"] == "flujo_13_semanas/v1"
    assert datos["corte"] == CORTE and datos["moneda"] == "MXN"
    semanas = datos["semanas"]
    assert [s["semana"] for s in semanas] == list(range(1, 14))
    assert semanas[0]["inicio"] == CORTE and semanas[-1]["fin"] == "2027-01-31"
    for previa, siguiente in zip(semanas, semanas[1:]):
        assert siguiente["saldo_inicial"] == previa["saldo_final"]
    for semana in semanas:
        montos = [semana["saldo_inicial"], semana["saldo_final"], semana["flujo_neto"],
                  semana["total_entradas"], semana["total_salidas"],
                  *semana["entradas"].values(), *semana["salidas"].values()]
        assert all(isinstance(m, str) and MONTO_TXT.match(m) for m in montos)
        neto = sum(D(m) for m in [*semana["entradas"].values(), *semana["salidas"].values()])
        assert D(semana["saldo_final"]) == D(semana["saldo_inicial"]) + neto
    assert datos["totales"]["saldo_final"] == semanas[-1]["saldo_final"]
    assert datos["totales"]["saldo_inicial"] == semanas[0]["saldo_inicial"]


def test_json_y_jsonl_son_deterministas(tmp_path: Path) -> None:
    assert _correr(tmp_path / "a") == 0
    assert _correr(tmp_path / "b") == 0
    for nombre in ("flujo_13_semanas.json", "filas_contrato.jsonl"):
        assert _sha(tmp_path / "a" / nombre) == _sha(tmp_path / "b" / nombre)
    texto = (tmp_path / "a" / "flujo_13_semanas.json").read_text(encoding="utf-8")
    assert "\r" not in texto and str(tmp_path) not in texto and str(RAIZ) not in texto


def test_cada_ruta_de_la_config_se_usa_y_queda_trazada(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    config = json.loads(CONFIG_EJEMPLO.read_text(encoding="utf-8"))
    fuentes = {f["fuente"]: f for f in _json(tmp_path)["fuentes"]}
    assert set(fuentes) == set(config["insumos"])
    for fuente, relativa in config["insumos"].items():
        ruta = (CONFIG_EJEMPLO.parent / relativa).resolve()
        assert fuentes[fuente]["archivo"] == ruta.name
        assert fuentes[fuente]["sha256"] == _sha(ruta)
        assert fuentes[fuente]["registros"] > 0


def test_cada_insumo_modificado_cambia_la_salida(tmp_path: Path) -> None:
    assert _correr(tmp_path / "base") == 0
    base = _sha(tmp_path / "base" / "flujo_13_semanas.json")
    config = json.loads(CONFIG_EJEMPLO.read_text(encoding="utf-8"))
    for fuente, relativa in config["insumos"].items():
        copia = tmp_path / f"datos_{fuente}"
        shutil.copytree(DATOS, copia)
        ruta = copia / Path(relativa).name
        lineas = ruta.read_text(encoding="utf-8").splitlines()
        lineas.pop()  # quita el último registro de esa fuente
        ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")
        insumos = {k: str(copia / Path(v).name) for k, v in config["insumos"].items()}
        ruta_config = _copiar_config(tmp_path, insumos=insumos)
        assert _correr(tmp_path / f"out_{fuente}", config=ruta_config) == 0
        assert _sha(tmp_path / f"out_{fuente}" / "flujo_13_semanas.json") != base, fuente


def test_totales_por_fuente_cuadran_con_el_oraculo_del_generador(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    oraculo = json.loads((DATOS / "totales_control.json").read_text(encoding="utf-8"))
    signo = {"saldos_bancarios": 1, "cxc_antiguedad": 1, "cxp_programadas": -1, "nomina": -1, "impuestos": -1}
    for fuente in _json(tmp_path)["fuentes"]:
        control = oraculo["fuentes"][fuente["fuente"]]
        assert fuente["registros"] == control["registros"]
        assert D(fuente["total_insumo"]) == signo[fuente["fuente"]] * D(control["suma"])
        partes = D(fuente["en_horizonte"]) + D(fuente["fuera_horizonte"]) + D(fuente["no_cobrable"])
        assert partes == D(fuente["total_insumo"])


def test_jsonl_suma_lo_mismo_que_la_tabla(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    lineas = (tmp_path / "filas_contrato.jsonl").read_text(encoding="utf-8").splitlines()
    filas = [json.loads(linea) for linea in lineas]
    claves = [(f["semana"], f["concepto"], f["fuente"], f["referencia"]) for f in filas]
    assert claves == sorted(claves)
    datos = _json(tmp_path)
    flujo_total = sum((D(f["monto"]) for f in filas if f["concepto"] != "saldo_inicial"), D("0"))
    assert D(datos["totales"]["flujo_neto"]) == flujo_total
    assert D(datos["totales"]["saldo_final"]) == D(datos["totales"]["saldo_inicial"]) + flujo_total


# --- xlsx -------------------------------------------------------------------------------


def test_xlsx_tiene_13_semanas_y_montos_numericos(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    libro = load_workbook(tmp_path / "flujo_13_semanas.xlsx")
    assert libro.sheetnames == ["Flujo", "Filas contrato", "Fuentes", "Resumen"]
    hoja = libro["Flujo"]
    encabezado = [c.value for c in hoja[1]]
    assert encabezado == ["Concepto"] + [f"Semana {n}" for n in range(1, 14)]
    filas_monto = [fila for fila in hoja.iter_rows(min_row=4) if fila[0].value]
    assert len(filas_monto) >= 9
    for fila in filas_monto:
        for celda in fila[1:]:
            assert isinstance(celda.value, (int, float)) and not isinstance(celda.value, bool), celda.coordinate
            assert celda.number_format == "#,##0.00"
    datos = _json(tmp_path)
    por_etiqueta = {fila[0].value: [c.value for c in fila[1:]] for fila in filas_monto}
    finales = [D(str(v)).quantize(D("0.01"), rounding=ROUND_HALF_UP) for v in por_etiqueta["Saldo final"]]
    assert finales == [D(s["saldo_final"]) for s in datos["semanas"]]


def test_xlsx_hojas_de_detalle_no_guardan_montos_como_texto(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    libro = load_workbook(tmp_path / "flujo_13_semanas.xlsx")
    for nombre in ("Filas contrato", "Fuentes", "Resumen"):
        hoja = libro[nombre]
        textos: List[str] = []
        for fila in hoja.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and MONTO_TXT.match(celda.value):
                    textos.append(celda.coordinate)
                if isinstance(celda.value, (int, float)) and celda.number_format == "#,##0.00":
                    assert not isinstance(celda.value, bool)
        assert textos == [], nombre
    filas = libro["Filas contrato"]
    assert filas.max_row - 1 == len((tmp_path / "filas_contrato.jsonl").read_text(encoding="utf-8").splitlines())


def test_xlsx_es_byte_a_byte_determinista_aunque_cambie_el_reloj(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    assert _correr(tmp_path / "a") == 0
    reloj_real = time.time
    monkeypatch.setattr(time, "time", lambda: reloj_real() + SEGUNDOS_POR_DIA)
    assert _correr(tmp_path / "b") == 0
    assert _sha(tmp_path / "a" / "flujo_13_semanas.xlsx") == _sha(tmp_path / "b" / "flujo_13_semanas.xlsx")


def test_xlsx_lleva_metadatos_y_fechas_de_zip_fijas(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    ruta = tmp_path / "flujo_13_semanas.xlsx"
    with zipfile.ZipFile(ruta) as archivo:
        entradas = archivo.infolist()
    assert entradas and all(e.date_time == FECHA_ZIP for e in entradas)
    propiedades = load_workbook(ruta).properties
    medianoche_corte = datetime.fromisoformat(f"{CORTE}T00:00:00")
    assert propiedades.creator == "flujo"
    assert propiedades.created == medianoche_corte
    assert propiedades.modified == medianoche_corte


def test_xlsx_filas_contrato_incluye_la_marca(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    hoja = load_workbook(tmp_path / "flujo_13_semanas.xlsx")["Filas contrato"]
    encabezado = [c.value for c in hoja[1]]
    assert encabezado[-1] == "marca"
    marcas = {fila[-1].value for fila in hoja.iter_rows(min_row=2)}
    assert {"vencido", "saldo_acreedor", None} <= marcas


# --- egresos vencidos, saldos acreedores y saldos en cero -------------------------------


def _filas_jsonl(salida: Path) -> List[Dict[str, Any]]:
    lineas = (salida / "filas_contrato.jsonl").read_text(encoding="utf-8").splitlines()
    return [json.loads(linea) for linea in lineas]


def test_avisos_cuadran_con_los_casos_plantados_por_el_generador(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
) -> None:
    assert _correr(tmp_path) == 0
    casos = json.loads((DATOS / "totales_control.json").read_text(encoding="utf-8"))["casos"]
    avisos = _json(tmp_path)["avisos"]
    assert set(avisos) == {"cxc_saldo_acreedor", "cxc_saldo_cero", "egresos_vencidos"}
    assert avisos["egresos_vencidos"]["registros"] == casos["egresos_vencidos"]["registros"] > 0
    assert D(avisos["egresos_vencidos"]["monto"]) == -D(casos["egresos_vencidos"]["suma"])
    assert avisos["cxc_saldo_acreedor"]["registros"] == casos["cxc_saldo_acreedor"]["registros"] > 0
    assert D(avisos["cxc_saldo_acreedor"]["monto"]) == D(casos["cxc_saldo_acreedor"]["suma"]) < 0
    assert avisos["cxc_saldo_cero"]["registros"] == casos["cxc_saldo_cero"]["registros"] > 0
    resumen = json.loads(capsys.readouterr().out)
    assert resumen["avisos"] == {clave: valor["registros"] for clave, valor in avisos.items()}
    cxc = {f["fuente"]: f for f in _json(tmp_path)["fuentes"]}["cxc_antiguedad"]
    assert cxc["registros_en_cero"] == casos["cxc_saldo_cero"]["registros"]


def test_filas_marcadas_van_en_la_semana_1_con_fecha_de_corte(tmp_path: Path) -> None:
    assert _correr(tmp_path) == 0
    filas = _filas_jsonl(tmp_path)
    assert {f["marca"] for f in filas} == {None, "vencido", "saldo_acreedor"}
    for fila in filas:
        if fila["marca"] is not None:
            assert (fila["semana"], fila["fecha"]) == (1, CORTE), fila
        if fila["marca"] == "vencido":
            assert fila["concepto"].startswith("pago_") and D(fila["monto"]) < 0
        if fila["marca"] == "saldo_acreedor":
            assert fila["concepto"] == "cobranza_clientes" and D(fila["monto"]) < 0
    vencidas = sum((D(f["monto"]) for f in filas if f["marca"] == "vencido"), D("0"))
    assert D(_json(tmp_path)["avisos"]["egresos_vencidos"]["monto"]) == vencidas


# --- errores y códigos de salida --------------------------------------------------------


def test_insumo_con_esquema_invalido_sale_con_3_y_no_escribe(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
) -> None:
    copia = tmp_path / "datos"
    shutil.copytree(DATOS, copia)
    ruta = copia / "cxp_programadas.csv"
    lineas = ruta.read_text(encoding="utf-8").splitlines()
    lineas[0] = lineas[0].replace("monto", "importe")
    ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")
    config = json.loads(CONFIG_EJEMPLO.read_text(encoding="utf-8"))
    insumos = {k: str(copia / Path(v).name) for k, v in config["insumos"].items()}
    salida = tmp_path / "out"
    assert _correr(salida, config=_copiar_config(tmp_path, insumos=insumos)) == 3
    assert not salida.exists()
    assert "cxp_programadas.csv" in capsys.readouterr().err


def test_corte_que_no_corresponde_a_los_saldos_sale_con_3(tmp_path: Path) -> None:
    assert _correr(tmp_path / "out", corte="2026-11-03") == 3
    assert not (tmp_path / "out").exists()


def test_insumo_inexistente_sale_con_3(tmp_path: Path) -> None:
    config = json.loads(CONFIG_EJEMPLO.read_text(encoding="utf-8"))
    insumos = {k: str((CONFIG_EJEMPLO.parent / v).resolve()) for k, v in config["insumos"].items()}
    insumos["nomina"] = str(tmp_path / "no_existe.csv")
    assert _correr(tmp_path / "out", config=_copiar_config(tmp_path, insumos=insumos)) == 3


@pytest.mark.parametrize("argumentos", [
    ["--corte", CORTE, "--salida", "x"],
    ["--config", "c.json", "--salida", "x"],
    ["--config", "c.json", "--corte", CORTE],
    ["--config", "c.json", "--corte", "02/11/2026", "--salida", "x"],
    ["--config", "c.json", "--corte", "2026-02-30", "--salida", "x"],
])
def test_argumentos_faltantes_o_invalidos_salen_con_2(argumentos: List[str]) -> None:
    with pytest.raises(SystemExit) as salida:
        main(argumentos)
    assert salida.value.code == 2


@pytest.mark.parametrize("cambios, mensaje", [
    ({"extra": 1}, "desconocidas"),
    ({"esquema": "otro/v9"}, "esquema"),
    ({"moneda": "pesos"}, "moneda"),
    ({"insumos": {"saldos_bancarios": "a.csv"}}, "faltan"),
    ({"curva_cobranza": []}, "al menos un tramo"),
])
def test_config_invalida_sale_con_2(tmp_path: Path, capsys: pytest.CaptureFixture[str],
                                    cambios: Dict[str, Any], mensaje: str) -> None:
    assert _correr(tmp_path / "out", config=_copiar_config(tmp_path, **cambios)) == 2
    assert mensaje in capsys.readouterr().err
    assert not (tmp_path / "out").exists()


def test_config_inexistente_o_no_json_sale_con_2(tmp_path: Path) -> None:
    assert _correr(tmp_path / "out", config=tmp_path / "no_existe.json") == 2
    mala = tmp_path / "mala.json"
    mala.write_text("{no es json", encoding="utf-8")
    assert _correr(tmp_path / "out", config=mala) == 2
    repetida = tmp_path / "repetida.json"
    repetida.write_text('{"esquema": "flujo_config/v1", "esquema": "flujo_config/v1"}', encoding="utf-8")
    assert _correr(tmp_path / "out", config=repetida) == 2
    lista = tmp_path / "lista.json"
    lista.write_text("[]", encoding="utf-8")
    assert _correr(tmp_path / "out", config=lista) == 2


def test_salida_que_es_un_archivo_sale_con_1(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    ocupada = tmp_path / "ocupada"
    ocupada.write_text("x", encoding="utf-8")
    assert _correr(ocupada) == 1
    assert "no se pudo escribir" in capsys.readouterr().err


# --- punta a punta ----------------------------------------------------------------------


def test_python_m_flujo_punta_a_punta(tmp_path: Path) -> None:
    entorno = {**os.environ, "PYTHONPATH": str(MODULO)}
    corrida = subprocess.run(
        [sys.executable, "-m", "flujo", "--config", str(CONFIG_EJEMPLO), "--corte", CORTE,
         "--salida", str(tmp_path)],
        capture_output=True, text=True, env=entorno, cwd=str(tmp_path), check=False,
    )
    assert corrida.returncode == 0, corrida.stderr
    resumen = json.loads(corrida.stdout)
    assert resumen["corte"] == CORTE
    assert resumen["saldo_final"] == _json(tmp_path)["totales"]["saldo_final"]
