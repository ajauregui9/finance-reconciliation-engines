"""Libro xlsx del cierre con montos numéricos, y su verificación releyendo el archivo.

La verificación no confía en el objeto en memoria que escribió el libro: abre el xlsx
guardado, convierte cada celda por la frontera de montos y la compara contra el resultado.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Callable, List, Optional, Sequence, Tuple, Union

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from .cierre import CierreMes, ResultadoCorrida
from .config import CONCEPTOS_TARIFA
from .dinero import CERO, a_decimal, a_tasa, texto_tasa
from .errores import ErrorInsumo
from .salida_json import seleccionar_meses

FORMATO_MONTO = "#,##0.00"
FORMATO_TASA = "0.000000"
HOJA_RESUMEN = "Resumen"
HOJA_GASTO = "Gasto devengado"
HOJA_ALERTAS = "Alertas"
HOJA_CANDADOS = "Candados"
HOJAS = (HOJA_RESUMEN, HOJA_GASTO, HOJA_ALERTAS, HOJA_CANDADOS)

Valor = Union[Decimal, int, str, None]
MONTO, TASA, TEXTO = "monto", "tasa", "texto"


@dataclass(frozen=True)
class _Columna:
    titulo: str
    tipo: str
    valor: Callable[[CierreMes], Valor]


def _banda(cierre: CierreMes) -> str:
    dentro = cierre.retencion.dentro_de_banda
    return "sin base" if dentro is None else ("sí" if dentro else "no")


COLUMNAS_RESUMEN: Tuple[_Columna, ...] = (
    _Columna("Mes", TEXTO, lambda c: c.mes),
    _Columna("Filas", TEXTO, lambda c: c.grupos.todas.filas),
    _Columna("Total liquidado", MONTO, lambda c: c.grupos.todas.total),
    _Columna("Venta neta cohorte del mes", MONTO, lambda c: c.grupos.cohorte_mes.venta_neta),
    _Columna("Venta neta cola", MONTO, lambda c: c.grupos.cola.venta_neta),
    _Columna("Venta ERP", MONTO, lambda c: c.venta_erp),
    _Columna("Venta pendiente", MONTO, lambda c: c.venta_pendiente),
    _Columna("Comisión", MONTO, lambda c: c.gasto.tarifas["comision"].total),
    _Columna("Logística", MONTO, lambda c: c.gasto.tarifas["logistica"].total),
    _Columna("Otras tarifas", MONTO, lambda c: c.gasto.tarifas["otras_tarifas"].total),
    _Columna("Cargos del periodo", MONTO, lambda c: sum(c.gasto.cargos_periodo.values(), CERO)),
    _Columna("Publicidad", MONTO, lambda c: c.gasto.publicidad),
    _Columna("Gasto devengado", MONTO, lambda c: c.gasto.total),
    _Columna("Retención retenida", MONTO, lambda c: c.retencion.retenida),
    _Columna("Retención esperada", MONTO, lambda c: c.retencion.esperada),
    _Columna("Tasa de retención observada", TASA, lambda c: c.retencion.tasa_observada),
    _Columna("Retención dentro de banda", TEXTO, _banda),
    _Columna("Candados", TEXTO, lambda c: "ok" if c.ok else "falla"),
)
ENCABEZADO_GASTO = ("Mes", "Concepto", "Componente", "Monto")


def lineas_gasto(cierre: CierreMes) -> List[Tuple[str, str, str, Decimal]]:
    """El gasto devengado en formato largo: una fila por componente."""
    lineas: List[Tuple[str, str, str, Decimal]] = []
    for concepto in CONCEPTOS_TARIFA:
        linea = cierre.gasto.tarifas[concepto]
        for componente, monto in (("cohorte_mes", linea.cohorte_mes), ("provision", linea.provision),
                                  ("cola_real", linea.cola_real), ("reverso", linea.reverso),
                                  ("total", linea.total)):
            lineas.append((cierre.mes, concepto, componente, monto))
    for tipo, monto in cierre.gasto.cargos_periodo.items():
        lineas.append((cierre.mes, "cargos_periodo", tipo, monto))
    lineas.append((cierre.mes, "publicidad", cierre.publicidad.fuente_autoritativa, cierre.gasto.publicidad))
    lineas.append((cierre.mes, "total", "gasto_devengado", cierre.gasto.total))
    return lineas


def _formato(tipo: str) -> Optional[str]:
    return {MONTO: FORMATO_MONTO, TASA: FORMATO_TASA}.get(tipo)


def _escribir_tabla(hoja: Worksheet, encabezado: Sequence[str], filas: Sequence[Sequence[Valor]],
                    tipos: Sequence[str]) -> None:
    hoja.append(list(encabezado))
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        hoja.append(list(fila))
        for celda, tipo in zip(hoja[hoja.max_row], tipos):
            formato = _formato(tipo)
            if formato is not None:
                celda.number_format = formato
    hoja.freeze_panes = "A2"


def _filas_resumen(cierres: Sequence[CierreMes]) -> List[List[Valor]]:
    return [[col.valor(c) for col in COLUMNAS_RESUMEN] for c in cierres]


def _filas_gasto(cierres: Sequence[CierreMes]) -> List[List[Valor]]:
    return [[mes, concepto, componente, monto] for c in cierres for mes, concepto, componente, monto in lineas_gasto(c)]


def escribir_xlsx(resultado: ResultadoCorrida, ruta: Path, meses: Optional[Sequence[str]] = None) -> None:
    """Escribe las cuatro hojas. Los montos van como números, nunca como texto."""
    cierres = seleccionar_meses(resultado, meses)
    libro = Workbook()
    resumen = libro.active
    if resumen is None:
        raise RuntimeError("openpyxl no creó la hoja inicial")
    resumen.title = HOJA_RESUMEN
    _escribir_tabla(resumen, [c.titulo for c in COLUMNAS_RESUMEN], _filas_resumen(cierres),
                    [c.tipo for c in COLUMNAS_RESUMEN])
    _escribir_tabla(libro.create_sheet(HOJA_GASTO), ENCABEZADO_GASTO, _filas_gasto(cierres),
                    (TEXTO, TEXTO, TEXTO, MONTO))
    alertas = [[a.mes or "corrida", a.codigo, a.mensaje] for a in resultado.alertas]
    alertas += [[c.mes, a.codigo, a.mensaje] for c in cierres for a in c.alertas]
    _escribir_tabla(libro.create_sheet(HOJA_ALERTAS), ("Mes", "Código", "Mensaje"), alertas, (TEXTO,) * 3)
    candados = [[c.mes, k.nombre, "sí" if k.ok else "no", k.detalle] for c in cierres for k in c.candados]
    _escribir_tabla(libro.create_sheet(HOJA_CANDADOS), ("Mes", "Candado", "OK", "Detalle"), candados, (TEXTO,) * 4)
    libro.save(ruta)


# ------------------------------------------------------------------ verificación


def _comparar_celda(lugar: str, tipo: str, leido: object, esperado: Valor) -> Optional[str]:
    if esperado is None or tipo == TEXTO:
        return None if leido == esperado else f"{lugar}: se leyó {leido!r} y se esperaba {esperado!r}"
    if isinstance(leido, str):
        return f"{lugar}: el monto está como texto ({leido!r})"
    if not isinstance(leido, (int, float, Decimal)) or isinstance(leido, bool):
        return f"{lugar}: se esperaba un número y se leyó {leido!r}"
    if not isinstance(esperado, Decimal):
        return f"{lugar}: valor esperado no numérico {esperado!r}"
    if tipo == TASA:
        leido_texto, esperado_texto = texto_tasa(a_tasa(str(leido))), texto_tasa(esperado)
        return None if leido_texto == esperado_texto else f"{lugar}: tasa {leido_texto} vs {esperado_texto}"
    leido_monto = a_decimal(leido)
    return None if leido_monto == esperado else f"{lugar}: monto {leido_monto} vs {esperado}"


def _comparar_hoja(hoja: Worksheet, esperadas: Sequence[Sequence[Valor]], tipos: Sequence[str]) -> List[str]:
    leidas = [list(fila) for fila in hoja.iter_rows(min_row=2, values_only=True)]
    diferencias = []
    if len(leidas) != len(esperadas):
        diferencias.append(f"{hoja.title}: {len(leidas)} filas y se esperaban {len(esperadas)}")
    for numero, (leida, esperada) in enumerate(zip(leidas, esperadas), start=2):
        for columna, (tipo, valor_leido, valor_esperado) in enumerate(zip(tipos, leida, esperada), start=1):
            diferencia = _comparar_celda(f"{hoja.title}!{hoja.cell(row=numero, column=columna).coordinate}",
                                         tipo, valor_leido, valor_esperado)
            if diferencia is not None:
                diferencias.append(diferencia)
    return diferencias


def verificar_xlsx(ruta: Path, resultado: ResultadoCorrida,
                   meses: Optional[Sequence[str]] = None) -> Tuple[str, ...]:
    """Relee el libro y devuelve las diferencias contra el resultado (vacío si cuadra)."""
    try:
        libro = load_workbook(ruta, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        raise ErrorInsumo(f"{ruta.name}: no se pudo abrir el xlsx ({exc})") from exc
    faltantes = [f"falta la hoja '{h}'" for h in HOJAS if h not in libro.sheetnames]
    if faltantes:
        return tuple(faltantes)
    cierres = seleccionar_meses(resultado, meses)
    diferencias = _comparar_hoja(libro[HOJA_RESUMEN], _filas_resumen(cierres), [c.tipo for c in COLUMNAS_RESUMEN])
    diferencias += _comparar_hoja(libro[HOJA_GASTO], _filas_gasto(cierres), (TEXTO, TEXTO, TEXTO, MONTO))
    return tuple(diferencias)
