"""Libro xlsx del flujo (openpyxl). Todo monto es una celda numérica con formato; nunca texto.

El libro se arma en memoria y se devuelve como bytes, así el CLI no escribe nada
hasta que todas las salidas están listas. Los bytes son deterministas: los metadatos
llevan la medianoche de la fecha de corte y el zip se empaca con fechas fijas.
"""
from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from typing import Optional, Sequence, Union

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .avisos import avisos
from .contrato import (
    CONCEPTO_COBRANZA,
    CONCEPTO_DE_FUENTE,
    CONCEPTO_IMPUESTOS,
    CONCEPTO_NOMINA,
    CONCEPTO_PROVEEDORES,
    CONCEPTOS_FLUJO,
)
from .empaque_xlsx import bytes_deterministas
from .ensamblador import Flujo

FORMATO_MONTO = "#,##0.00"
FORMATO_FECHA = "yyyy-mm-dd"
ANCHO_ETIQUETA = 24
ANCHO_COLUMNA = 16
FILA_ENCABEZADO = 1
FILA_INICIO = 2
FILA_FIN = 3
PRIMERA_FILA_DATOS = 2
COLUMNA_PRIMER_MONTO = 2
HOJA_FLUJO = "Flujo"
HOJA_FILAS = "Filas contrato"
HOJA_FUENTES = "Fuentes"
HOJA_RESUMEN = "Resumen"
FILA_SEPARACION_AVISOS = 2
ETIQUETAS = {
    CONCEPTO_COBRANZA: "Cobranza clientes",
    CONCEPTO_PROVEEDORES: "Pago proveedores",
    CONCEPTO_NOMINA: "Pago nómina",
    CONCEPTO_IMPUESTOS: "Pago impuestos",
}

Valor = Optional[Union[str, int, date, Decimal]]


def _escribir_fila(hoja: Worksheet, numero: int, valores: Sequence[Valor], negrita: bool = False) -> None:
    """Escribe una fila; los `Decimal` llevan formato de monto y las fechas formato ISO."""
    for columna, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=numero, column=columna, value=valor)
        if isinstance(valor, Decimal):
            celda.number_format = FORMATO_MONTO
        elif isinstance(valor, date):
            celda.number_format = FORMATO_FECHA
        if negrita:
            celda.font = Font(bold=True)


def _hoja_flujo(hoja: Worksheet, flujo: Flujo) -> None:
    semanas = flujo.semanas
    _escribir_fila(hoja, FILA_ENCABEZADO, ["Concepto"] + [f"Semana {s.numero}" for s in semanas], negrita=True)
    _escribir_fila(hoja, FILA_INICIO, ["Inicio"] + [s.inicio for s in semanas])
    _escribir_fila(hoja, FILA_FIN, ["Fin"] + [s.fin for s in semanas])
    renglones = (
        [("Saldo inicial", [s.saldo_inicial for s in semanas])]
        + [(ETIQUETAS[c], [dict(s.montos)[c] for s in semanas]) for c in CONCEPTOS_FLUJO]
        + [
            ("Total entradas", [s.entradas for s in semanas]),
            ("Total salidas", [s.salidas for s in semanas]),
            ("Flujo neto", [s.flujo_neto for s in semanas]),
            ("Saldo final", [s.saldo_final for s in semanas]),
        ]
    )
    primera = hoja.max_row + 1
    for desfase, (etiqueta, montos) in enumerate(renglones):
        _escribir_fila(hoja, primera + desfase, [etiqueta, *montos], negrita=etiqueta == "Saldo final")
    hoja.freeze_panes = f"{get_column_letter(COLUMNA_PRIMER_MONTO)}{primera}"


def _hoja_filas(hoja: Worksheet, flujo: Flujo) -> None:
    columnas = ["fuente", "concepto", "semana", "fecha", "referencia", "contraparte", "monto", "marca"]
    _escribir_fila(hoja, FILA_ENCABEZADO, columnas, negrita=True)
    for numero, f in enumerate(flujo.filas, start=PRIMERA_FILA_DATOS):
        _escribir_fila(hoja, numero, [
            f.fuente, f.concepto, f.semana, f.fecha, f.referencia, f.contraparte, f.monto, f.marca,
        ])


def _hoja_fuentes(hoja: Worksheet, flujo: Flujo) -> None:
    columnas = ["fuente", "concepto", "archivo", "registros", "registros_en_cero", "filas_contrato",
                "total_insumo", "en_horizonte", "fuera_horizonte", "no_cobrable", "sha256"]
    _escribir_fila(hoja, FILA_ENCABEZADO, columnas, negrita=True)
    for numero, r in enumerate(flujo.fuentes, start=PRIMERA_FILA_DATOS):
        _escribir_fila(hoja, numero, [
            r.fuente, CONCEPTO_DE_FUENTE[r.fuente], r.archivo, r.registros, r.registros_en_cero, len(r.filas),
            r.total_insumo, r.en_horizonte, r.fuera_horizonte, r.no_cobrable, r.sha256,
        ])


def _hoja_resumen(hoja: Worksheet, flujo: Flujo) -> None:
    renglones: Sequence[Sequence[Valor]] = (
        ("Corte", flujo.corte),
        ("Moneda", flujo.moneda),
        ("Fin del horizonte", flujo.semanas[-1].fin),
        ("Semanas", len(flujo.semanas)),
        ("Saldo inicial", flujo.semanas[0].saldo_inicial),
        ("Total entradas", flujo.total_entradas),
        ("Total salidas", flujo.total_salidas),
        ("Saldo final", flujo.semanas[-1].saldo_final),
    )
    for numero, renglon in enumerate(renglones, start=1):
        _escribir_fila(hoja, numero, renglon)
    primera = hoja.max_row + FILA_SEPARACION_AVISOS
    _escribir_fila(hoja, primera, ["Aviso", "Registros", "Monto"], negrita=True)
    for desfase, aviso in enumerate(avisos(flujo), start=1):
        _escribir_fila(hoja, primera + desfase, [aviso.descripcion, aviso.registros, aviso.monto])


def _anchos(hoja: Worksheet) -> None:
    for indice in range(1, hoja.max_column + 1):
        ancho = ANCHO_ETIQUETA if indice == 1 else ANCHO_COLUMNA
        hoja.column_dimensions[get_column_letter(indice)].width = ancho


def libro_xlsx(flujo: Flujo) -> bytes:
    libro = Workbook()
    hoja_flujo = libro.active
    if not isinstance(hoja_flujo, Worksheet):
        raise TypeError("openpyxl no creó una hoja de cálculo activa")
    hoja_flujo.title = HOJA_FLUJO
    _hoja_flujo(hoja_flujo, flujo)
    _hoja_filas(libro.create_sheet(HOJA_FILAS), flujo)
    _hoja_fuentes(libro.create_sheet(HOJA_FUENTES), flujo)
    _hoja_resumen(libro.create_sheet(HOJA_RESUMEN), flujo)
    for hoja in libro.worksheets:
        _anchos(hoja)
    return bytes_deterministas(libro, datetime.combine(flujo.corte, time.min))
